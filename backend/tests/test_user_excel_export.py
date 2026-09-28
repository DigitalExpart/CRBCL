"""Tests for CRBCL User Directory Excel Export — corrected contract.

Verifies:
 1. Authorization: Unauthenticated (401), Unauthorized (403), Authorized Admin (200).
 2. Exact column order and headings:
        Username | First Name | Last Name | Status | Email | Last Login
    Column 6 MUST be "Last Login", NOT "Account Created".
 3. Active/Inactive human-readable conversion.
 4. Column 6 uses User.last_login_at — the authoritative successful-login timestamp.
    It MUST NOT use created_at, updated_at, token-refresh time, or export time.
 5. Never-logged-in users produce a blank (None) cell in column 6.
 6. Failed authentication does NOT change last_login_at.
 7. Successful password-login DOES update last_login_at.
 8. Successful OTP-verify login DOES update last_login_at.
 9. Special characters (Cree syllabics, accents, hyphens, quotes).
10. Blank optional fields (single name, empty name).
11. Formula-injection protection (=, +, -, @, \\t, \\r).
12. Workbook validity (freeze panes A2, autofilter, column widths).
13. Response filename and Content-Type.
14. Soft-deleted users are excluded from the export.
15. Inactive users appear with "Inactive" status.
16. Username = User.email (no separate username column; email is the auth identity).
    First Name / Last Name = derived from User.full_name split — NOT authoritative.
17. Successful-login bookkeeping is consolidated into
    AuthService.record_successful_login(); the router does NOT write
    last_login_at directly.
"""

from __future__ import annotations

import io
import uuid
from datetime import UTC, date, datetime, timedelta

import openpyxl
import pytest
from httpx import AsyncClient
from sqlalchemy import select
from sqlalchemy.ext.asyncio import AsyncSession

from app.auth.security import hash_password
from app.auth.service import AuthService
from app.models.user import User
from app.services.user_export_service import (
    USER_EXPORT_COLUMNS,
    extract_user_row,
    generate_user_excel_export,
    sanitize_cell_value,
)

# ── Constants ──────────────────────────────────────────────────────────────────

EXPECTED_COLUMNS = ["Username", "First Name", "Last Name", "Status", "Email", "Last Login"]
FORBIDDEN_COLUMNS = ["Account Created", "Role", "Department", "permissions", "Phone", "Employee ID"]

# ── Unit Tests: Service Layer ──────────────────────────────────────────────────


def test_export_columns_constant_exact():
    """USER_EXPORT_COLUMNS must exactly match the approved six-column contract."""
    assert USER_EXPORT_COLUMNS == EXPECTED_COLUMNS
    assert len(USER_EXPORT_COLUMNS) == 6
    assert "Last Login" in USER_EXPORT_COLUMNS
    assert "Account Created" not in USER_EXPORT_COLUMNS, (
        "Column 6 must be 'Last Login', not 'Account Created'. "
        "Do not use created_at as a substitute."
    )


def test_sanitize_cell_value_formula_injection():
    """Verify formula injection triggers are escaped with leading single quote."""
    # Triggers: =, +, -, @, \t, \r
    assert sanitize_cell_value("=cmd|'/C calc'!A0") == "'=cmd|'/C calc'!A0"
    assert sanitize_cell_value("+12345") == "'+12345"
    assert sanitize_cell_value("-2+3") == "'-2+3"
    assert sanitize_cell_value("@SUM(A1:A10)") == "'@SUM(A1:A10)"
    assert sanitize_cell_value("\t=1+1") == "'\t=1+1"
    assert sanitize_cell_value("\r=cmd") == "'\r=cmd"

    # Leading whitespace before trigger
    assert sanitize_cell_value("  =1+1") == "'  =1+1"
    assert sanitize_cell_value("   +test") == "'   +test"

    # Normal strings should NOT be modified
    assert sanitize_cell_value("Jane Doe") == "Jane Doe"
    assert sanitize_cell_value("jane.doe@crbcl.ca") == "jane.doe@crbcl.ca"
    assert sanitize_cell_value("user+tag@crbcl.ca") == "user+tag@crbcl.ca"
    assert sanitize_cell_value("") == ""
    assert sanitize_cell_value(None) is None
    assert sanitize_cell_value(123) == 123


def _make_user(
    email: str,
    full_name: str,
    is_active: bool = True,
    last_login_at: datetime | None = None,
    created_at: datetime | None = None,
) -> User:
    """Helper that constructs a User without a DB session."""
    user = User(
        email=email,
        email_normalized=email.lower(),
        password_hash=hash_password("Password123!"),
        full_name=full_name,
        is_active=is_active,
    )
    # Simulate DB timestamps
    user.last_login_at = last_login_at
    user.created_at = created_at or datetime(2025, 1, 1, 0, 0, 0, tzinfo=UTC)
    return user


# ── Column 6 contract: Last Login, not Account Created ────────────────────────


def test_extract_user_row_uses_last_login_not_created_at():
    """Column 6 must be last_login_at; created_at MUST NOT appear as Last Login."""
    login_ts = datetime(2026, 6, 15, 9, 30, 0, tzinfo=UTC)
    created_ts = datetime(2025, 1, 1, 8, 0, 0, tzinfo=UTC)  # deliberately different

    user = _make_user(
        email="synthetic.login@crbcl.ca",
        full_name="Test Worker",
        last_login_at=login_ts,
        created_at=created_ts,
    )

    row = extract_user_row(user)

    # Column 6 key must be "Last Login"
    assert "Last Login" in row, "extract_user_row must return a 'Last Login' key"
    assert "Account Created" not in row, (
        "extract_user_row must NOT return 'Account Created'. "
        "The contract changed to Last Login."
    )

    # Verify it is last_login_at, not created_at
    expected_last_login = login_ts.astimezone(UTC).replace(tzinfo=None)
    assert row["Last Login"] == expected_last_login, (
        f"Last Login should be {expected_last_login!r} (from last_login_at), "
        f"got {row['Last Login']!r}. Do not substitute created_at."
    )

    # Guard: the two timestamps are different, so if we see created_at it's a bug
    created_naive = created_ts.astimezone(UTC).replace(tzinfo=None)
    assert row["Last Login"] != created_naive, (
        "Last Login returned created_at. This is the wrong business field."
    )


def test_extract_user_row_never_logged_in_blank():
    """User who has never logged in must produce None (blank cell) for Last Login."""
    user = _make_user(
        email="synthetic.neverlogin@crbcl.ca",
        full_name="Fresh Account",
        last_login_at=None,  # never logged in
        created_at=datetime(2026, 1, 1, tzinfo=UTC),
    )

    row = extract_user_row(user)

    assert row["Last Login"] is None, (
        "Never-logged-in users must have Last Login = None (blank cell). "
        "Do not substitute created_at or any other timestamp."
    )


def test_extract_user_row_exact_mapping():
    """Verify exact column mapping and Active/Inactive status for a logged-in user."""
    login_ts = datetime(2026, 3, 20, 10, 15, 0, tzinfo=UTC)
    user = _make_user(
        email="synthetic.worker@crbcl.ca",
        full_name="Sarah Connor",
        is_active=True,
        last_login_at=login_ts,
    )

    row = extract_user_row(user)

    assert row["Username"] == "synthetic.worker@crbcl.ca"
    assert row["First Name"] == "Sarah"
    assert row["Last Name"] == "Connor"
    assert row["Status"] == "Active"
    assert row["Email"] == "synthetic.worker@crbcl.ca"
    # Last Login is naive UTC datetime
    assert row["Last Login"] == datetime(2026, 3, 20, 10, 15, 0)


def test_extract_user_row_inactive_and_single_name():
    """Verify inactive user conversion and single-word name handling."""
    login_ts = datetime(2025, 11, 5, 14, 15, 0, tzinfo=UTC)
    user = _make_user(
        email="synthetic.inactive@crbcl.ca",
        full_name="Madonna",
        is_active=False,
        last_login_at=login_ts,
    )

    row = extract_user_row(user)

    assert row["Username"] == "synthetic.inactive@crbcl.ca"
    assert row["First Name"] == "Madonna"
    assert row["Last Name"] == ""
    assert row["Status"] == "Inactive"
    assert row["Email"] == "synthetic.inactive@crbcl.ca"
    assert row["Last Login"] == datetime(2025, 11, 5, 14, 15, 0)


def test_extract_user_row_username_is_email():
    """Username column must be User.email — no separate username column exists on User."""
    user = _make_user(
        email="specific.user@crbcl.ca",
        full_name="Specific User",
        last_login_at=datetime(2026, 1, 1, tzinfo=UTC),
    )
    row = extract_user_row(user)
    assert row["Username"] == "specific.user@crbcl.ca", (
        "Username must be User.email. No separate username column exists."
    )
    # Email and Username are the same field (email is the auth identity)
    assert row["Username"] == row["Email"]


def test_extract_user_row_names_are_derived_from_full_name():
    """First Name / Last Name are derived from full_name — NOT authoritative separate fields."""
    user = _make_user(
        email="derived.names@crbcl.ca",
        full_name="Alexandra Fontaine-Leblanc",
        last_login_at=None,
    )
    row = extract_user_row(user)
    # Derived by splitting on first whitespace only
    assert row["First Name"] == "Alexandra"
    assert row["Last Name"] == "Fontaine-Leblanc"
    # Verify these are NOT coming from non-existent model attributes
    assert not hasattr(user, "first_name"), (
        "User model must not have a first_name column — it stores only full_name."
    )
    assert not hasattr(user, "last_name"), (
        "User model must not have a last_name column — it stores only full_name."
    )


def test_extract_user_row_blank_full_name():
    """Verify blank optional fields produce empty strings, not None or errors."""
    user = _make_user(
        email="synthetic.noname@crbcl.ca",
        full_name="",
        last_login_at=datetime(2026, 4, 1, 12, 0, 0, tzinfo=UTC),
    )

    row = extract_user_row(user)

    assert row["First Name"] == ""
    assert row["Last Name"] == ""
    assert row["Status"] == "Active"
    # Last Login is still populated from last_login_at
    assert row["Last Login"] == datetime(2026, 4, 1, 12, 0, 0)


def test_generate_user_excel_export_column_contract():
    """Verify workbook column order, headings, freeze panes, autofilter, and widths.

    This is the canonical test for the six-column contract.
    """
    login_ts1 = datetime(2026, 2, 1, 10, 0, 0, tzinfo=UTC)
    login_ts2 = datetime(2025, 8, 15, 16, 45, 30, tzinfo=UTC)

    u1 = _make_user(
        email="synthetic.u1@crbcl.ca",
        full_name="Alice Smith",
        is_active=True,
        last_login_at=login_ts1,
    )
    u2 = _make_user(
        email="synthetic.u2@crbcl.ca",
        full_name="Bob Jones",
        is_active=False,
        last_login_at=login_ts2,
    )

    xlsx_bytes = generate_user_excel_export([u1, u2])

    wb = openpyxl.load_workbook(io.BytesIO(xlsx_bytes))
    ws = wb.active

    # Exact column order & headings
    headers = [cell.value for cell in ws[1]]
    assert headers == USER_EXPORT_COLUMNS
    assert headers == EXPECTED_COLUMNS
    assert headers[5] == "Last Login", f"Column 6 must be 'Last Login', got {headers[5]!r}"
    assert "Account Created" not in headers, (
        "'Account Created' must not appear in the export. Contract is 'Last Login'."
    )
    assert len(headers) == 6

    # No extra forbidden fields
    for forbidden in FORBIDDEN_COLUMNS:
        assert forbidden not in headers

    # Verify row count
    assert ws.max_row == 3  # header + 2 users

    # Verify user 1 — column-by-column
    assert ws.cell(2, 1).value == "synthetic.u1@crbcl.ca"   # Username
    assert ws.cell(2, 2).value == "Alice"                    # First Name
    assert ws.cell(2, 3).value == "Smith"                    # Last Name
    assert ws.cell(2, 4).value == "Active"                   # Status
    assert ws.cell(2, 5).value == "synthetic.u1@crbcl.ca"   # Email
    # Column 6: Last Login — must be last_login_at, not created_at
    assert ws.cell(2, 6).value == datetime(2026, 2, 1, 10, 0, 0)
    assert ws.cell(2, 6).number_format == "yyyy-mm-dd hh:mm:ss"

    # Verify user 2 (inactive)
    assert ws.cell(3, 4).value == "Inactive"
    assert ws.cell(3, 6).value == datetime(2025, 8, 15, 16, 45, 30)

    # Freeze panes at A2
    assert ws.freeze_panes == "A2"

    # Autofilter covers all columns and rows
    assert ws.auto_filter.ref == "A1:F3"

    # Column widths are set and readable (>= 14)
    for col_letter in ["A", "B", "C", "D", "E", "F"]:
        assert ws.column_dimensions[col_letter].width >= 14


def test_generate_user_excel_never_logged_in_blank_cell():
    """Never-logged-in users must have a blank (None) cell in column 6, not created_at."""
    u_never = _make_user(
        email="synthetic.nologin@crbcl.ca",
        full_name="Brand New",
        last_login_at=None,
        created_at=datetime(2026, 9, 1, tzinfo=UTC),  # has a created_at — must NOT appear
    )

    xlsx_bytes = generate_user_excel_export([u_never])
    wb = openpyxl.load_workbook(io.BytesIO(xlsx_bytes))
    ws = wb.active

    last_login_cell = ws.cell(2, 6)
    assert last_login_cell.value is None, (
        f"Never-logged-in user must have blank Last Login, got {last_login_cell.value!r}. "
        "Do not substitute created_at."
    )


# ── Integration Tests: HTTP API Endpoint ───────────────────────────────────────


@pytest.mark.asyncio
async def test_user_export_requires_auth(client: AsyncClient):
    """Anonymous request without auth token is rejected with 401 Unauthorized."""
    response = await client.get("/api/v1/users/export")
    assert response.status_code == 401


@pytest.mark.asyncio
async def test_user_export_forbidden_for_unauthorized_user(client: AsyncClient, caseworker_user):
    """Caseworker lacking ADMIN_USERS_MANAGE is rejected with 403 Forbidden."""
    response = await client.get("/api/v1/users/export", headers=caseworker_user["headers"])
    assert response.status_code == 403
    assert response.json()["error"]["code"] == "PERMISSION_DENIED"


@pytest.mark.asyncio
async def test_user_export_success_column_contract(
    client: AsyncClient,
    it_admin_user,
    db_session: AsyncSession,
):
    """IT Administrator export: verify exact six-column order and Last Login semantics.

    Checks:
    - Column 6 header is "Last Login" (not "Account Created")
    - Users with last_login_at have that value in column 6
    - Users who never logged in have blank in column 6
    - created_at is NEVER used as a substitute for Last Login
    """
    base_time = datetime(2026, 1, 1, 12, 0, 0, tzinfo=UTC)
    login_time = base_time + timedelta(days=10)  # deliberately later than created_at

    uid = uuid.uuid4().hex[:8]

    # User who HAS logged in
    u_logged = User(
        email=f"synthetic.logged.{uid}@crbcl.ca",
        email_normalized=f"synthetic.logged.{uid}@crbcl.ca",
        password_hash=hash_password("Secret123!"),
        full_name="Tânisi O'Connor-Smith",
        is_active=True,
    )
    u_logged.created_at = base_time
    u_logged.last_login_at = login_time

    # User who has NEVER logged in
    u_nologin = User(
        email=f"synthetic.nologin.{uid}@crbcl.ca",
        email_normalized=f"synthetic.nologin.{uid}@crbcl.ca",
        password_hash=hash_password("Secret123!"),
        full_name="René Gagnon",
        is_active=True,
    )
    u_nologin.created_at = base_time + timedelta(days=1)
    u_nologin.last_login_at = None

    # Inactive user
    u_inactive = User(
        email=f"synthetic.inact.{uid}@crbcl.ca",
        email_normalized=f"synthetic.inact.{uid}@crbcl.ca",
        password_hash=hash_password("Secret123!"),
        full_name="Inactive User",
        is_active=False,
    )
    u_inactive.created_at = base_time + timedelta(days=2)
    u_inactive.last_login_at = base_time + timedelta(days=5)

    db_session.add_all([u_logged, u_nologin, u_inactive])
    await db_session.commit()

    response = await client.get("/api/v1/users/export", headers=it_admin_user["headers"])
    assert response.status_code == 200

    # Content-Type and Filename
    content_type = response.headers.get("content-type", "")
    assert "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet" in content_type

    disposition = response.headers.get("content-disposition", "")
    today_str = date.today().isoformat()
    assert f'filename="crbcl-users-export-{today_str}.xlsx"' in disposition

    # Parse Excel workbook
    wb = openpyxl.load_workbook(io.BytesIO(response.content))
    ws = wb.active

    # Exact column headers
    headers = [cell.value for cell in ws[1]]
    assert headers == EXPECTED_COLUMNS, (
        f"Expected headers {EXPECTED_COLUMNS!r}, got {headers!r}. "
        "Column 6 must be 'Last Login', not 'Account Created'."
    )
    assert headers[5] == "Last Login"
    assert "Account Created" not in headers

    # Map rows by email to verify content
    exported_rows: dict[str, dict] = {}
    for r in range(2, ws.max_row + 1):
        email_val = ws.cell(r, 5).value
        if email_val:
            exported_rows[email_val] = {
                "username": ws.cell(r, 1).value,
                "first_name": ws.cell(r, 2).value,
                "last_name": ws.cell(r, 3).value,
                "status": ws.cell(r, 4).value,
                "email": email_val,
                "last_login": ws.cell(r, 6).value,
                "last_login_fmt": ws.cell(r, 6).number_format,
            }

    # 1. Logged-in user: last_login_at must appear in column 6, not created_at
    assert u_logged.email in exported_rows
    logged_row = exported_rows[u_logged.email]
    expected_login_naive = login_time.astimezone(UTC).replace(tzinfo=None)
    created_naive = base_time.astimezone(UTC).replace(tzinfo=None)
    assert logged_row["last_login"] is not None
    assert isinstance(logged_row["last_login"], datetime | date)
    assert logged_row["last_login"] == expected_login_naive, (
        f"Last Login should be {expected_login_naive!r} (last_login_at), "
        f"got {logged_row['last_login']!r}. Do not use created_at."
    )
    assert logged_row["last_login"] != created_naive, (
        "Last Login must NOT equal created_at. They are different business fields."
    )
    assert logged_row["last_login_fmt"] == "yyyy-mm-dd hh:mm:ss"

    # 2. Never-logged-in user: column 6 must be blank
    assert u_nologin.email in exported_rows
    nologin_row = exported_rows[u_nologin.email]
    assert nologin_row["last_login"] is None, (
        f"Never-logged-in user must have blank Last Login, got {nologin_row['last_login']!r}. "
        "Do not substitute created_at."
    )

    # 3. Inactive user: status must be "Inactive", last_login still populated
    assert u_inactive.email in exported_rows
    inact_row = exported_rows[u_inactive.email]
    assert inact_row["status"] == "Inactive"
    assert inact_row["last_login"] is not None

    # 4. Freeze panes and Autofilter
    assert ws.freeze_panes == "A2"
    assert ws.auto_filter.ref is not None
    assert ws.auto_filter.ref.startswith("A1:F")


@pytest.mark.asyncio
async def test_user_export_formula_injection_protection(
    client: AsyncClient,
    it_admin_user,
    db_session: AsyncSession,
):
    """Formula injection in names must be neutralised with a leading single quote."""
    uid = uuid.uuid4().hex[:8]
    u_injection = User(
        email=f"synthetic.inj.{uid}@crbcl.ca",
        email_normalized=f"synthetic.inj.{uid}@crbcl.ca",
        password_hash=hash_password("Secret123!"),
        full_name="=cmd|'/C calc'!A0 Payload",
        is_active=True,
    )
    u_injection.last_login_at = datetime(2026, 5, 1, tzinfo=UTC)
    db_session.add(u_injection)
    await db_session.commit()

    response = await client.get("/api/v1/users/export", headers=it_admin_user["headers"])
    assert response.status_code == 200

    wb = openpyxl.load_workbook(io.BytesIO(response.content))
    ws = wb.active

    for r in range(2, ws.max_row + 1):
        if ws.cell(r, 5).value == u_injection.email:
            cell_val = ws.cell(r, 2).value  # First Name
            assert cell_val is not None
            assert cell_val.startswith("'="), (
                f"Formula injection must be escaped with leading quote, got {cell_val!r}"
            )
            assert ws.cell(r, 2).data_type == "s"
            break
    else:
        pytest.fail("Injection test user not found in export")


@pytest.mark.asyncio
async def test_user_export_soft_deleted_excluded(
    client: AsyncClient,
    it_admin_user,
    db_session: AsyncSession,
):
    """Soft-deleted users must not appear in the export."""
    uid = uuid.uuid4().hex[:8]
    u_deleted = User(
        email=f"synthetic.deleted.{uid}@crbcl.ca",
        email_normalized=f"synthetic.deleted.{uid}@crbcl.ca",
        password_hash=hash_password("Secret123!"),
        full_name="Deleted Person",
        is_active=True,
    )
    u_deleted.last_login_at = datetime(2026, 1, 1, tzinfo=UTC)
    db_session.add(u_deleted)
    await db_session.flush()

    # Soft-delete via the deleted_at column only.
    # is_deleted is a read-only @property; setting deleted_at is sufficient.
    u_deleted.deleted_at = datetime.now(UTC)
    await db_session.commit()

    response = await client.get("/api/v1/users/export", headers=it_admin_user["headers"])
    assert response.status_code == 200

    wb = openpyxl.load_workbook(io.BytesIO(response.content))
    ws = wb.active

    emails_in_export = {ws.cell(r, 5).value for r in range(2, ws.max_row + 1)}
    assert u_deleted.email not in emails_in_export, (
        "Soft-deleted user must not appear in the export."
    )


# ── Authentication regression: last_login_at persistence ──────────────────────
# Bookkeeping is consolidated in AuthService.record_successful_login().
# The router must NOT write last_login_at directly.


@pytest.mark.asyncio
async def test_record_successful_login_service_method(db_session: AsyncSession):
    """AuthService.record_successful_login() is the single authoritative source.

    - Writes last_login_at to now.
    - Resets failed_login_count to 0.
    - Clears locked_until.
    - Does NOT flush/commit — caller is responsible.
    """
    from datetime import UTC

    from app.models.user import User as UserModel

    uid = uuid.uuid4().hex[:8]
    user = UserModel(
        email=f"synthetic.bookkeeping.{uid}@crbcl.ca",
        email_normalized=f"synthetic.bookkeeping.{uid}@crbcl.ca",
        password_hash=hash_password("Secret123!"),
        full_name="Bookkeeping Test",
        is_active=True,
        is_verified=True,
    )
    user.last_login_at = None
    user.failed_login_count = 3
    user.locked_until = datetime.now(UTC) + timedelta(minutes=5)
    db_session.add(user)
    await db_session.flush()

    svc = AuthService(db_session)
    before = datetime.now(UTC)
    await svc.record_successful_login(user)
    after = datetime.now(UTC)

    assert user.last_login_at is not None, (
        "record_successful_login must set last_login_at."
    )
    ts = user.last_login_at
    if ts.tzinfo is None:
        ts = ts.replace(tzinfo=UTC)
    assert before <= ts <= after
    assert user.failed_login_count == 0, (
        "record_successful_login must reset failed_login_count to 0."
    )
    assert user.locked_until is None, (
        "record_successful_login must clear locked_until."
    )



@pytest.mark.asyncio
async def test_successful_login_updates_last_login_at(client: AsyncClient, db_session: AsyncSession):
    """Successful password login must update last_login_at on the User record."""
    uid = uuid.uuid4().hex[:8]
    user = User(
        email=f"synthetic.logintest.{uid}@crbcl.ca",
        email_normalized=f"synthetic.logintest.{uid}@crbcl.ca",
        password_hash=hash_password("Secret123!"),
        full_name="Login Test User",
        is_active=True,
        is_verified=True,
    )
    user.last_login_at = None  # never logged in
    db_session.add(user)
    await db_session.commit()

    before = datetime.now(UTC)

    response = await client.post(
        "/api/v1/auth/login",
        json={"email": user.email, "password": "Secret123!"},
    )
    assert response.status_code == 200

    after = datetime.now(UTC)

    # Re-fetch from DB to verify persistence
    await db_session.refresh(user)
    assert user.last_login_at is not None, (
        "Successful login must persist last_login_at on the User record."
    )
    login_ts = user.last_login_at
    if login_ts.tzinfo is None:
        login_ts = login_ts.replace(tzinfo=UTC)
    assert before <= login_ts <= after, (
        f"last_login_at {login_ts!r} should be between {before!r} and {after!r}"
    )


@pytest.mark.asyncio
async def test_failed_login_does_not_update_last_login_at(client: AsyncClient, db_session: AsyncSession):
    """Failed authentication must NOT change last_login_at."""
    uid = uuid.uuid4().hex[:8]
    fixed_login = datetime(2026, 1, 15, 9, 0, 0, tzinfo=UTC)
    user = User(
        email=f"synthetic.failtest.{uid}@crbcl.ca",
        email_normalized=f"synthetic.failtest.{uid}@crbcl.ca",
        password_hash=hash_password("Secret123!"),
        full_name="Fail Test User",
        is_active=True,
        is_verified=True,
    )
    user.last_login_at = fixed_login
    db_session.add(user)
    await db_session.commit()

    response = await client.post(
        "/api/v1/auth/login",
        json={"email": user.email, "password": "WrongPassword!"},
    )
    assert response.status_code == 401

    # Reload and check last_login_at is unchanged
    await db_session.refresh(user)
    assert user.last_login_at is not None
    stored = user.last_login_at
    if stored.tzinfo is None:
        stored = stored.replace(tzinfo=UTC)
    assert stored == fixed_login, (
        f"Failed login must not change last_login_at. "
        f"Expected {fixed_login!r}, got {stored!r}."
    )


@pytest.mark.asyncio
async def test_otp_verify_login_updates_last_login_at(client: AsyncClient, db_session: AsyncSession):
    """OTP verify-login is a genuine login and must update last_login_at.

    The router delegates to AuthService.record_successful_login() —
    it must NOT write last_login_at directly.
    """
    # Register
    otp_email = f"synthetic.otplogin.{uuid.uuid4().hex[:8]}@crbcl.ca"
    reg_res = await client.post(
        "/api/v1/auth/register",
        json={"email": otp_email, "password": "Password123!", "full_name": "OTP Tester"},
    )
    assert reg_res.status_code == 200

    # Fetch user and confirm last_login_at is still None pre-OTP
    user_res = await db_session.execute(
        select(User).where(User.email_normalized == otp_email.lower())
    )
    user = user_res.scalar_one_or_none()
    assert user is not None
    assert user.last_login_at is None, "Freshly registered user must not have last_login_at set"

    # Generate a valid OTP code
    from app.services.email_service import EmailService

    svc = EmailService(db_session)
    valid_code = await svc.create_and_send_verification_code(otp_email)
    await db_session.commit()

    before = datetime.now(UTC)

    verify_res = await client.post(
        "/api/v1/auth/verify-otp",
        json={"email": otp_email, "otp_code": valid_code},
    )
    assert verify_res.status_code == 200

    after = datetime.now(UTC)

    # Reload user and confirm last_login_at was set
    await db_session.refresh(user)
    assert user.last_login_at is not None, (
        "OTP verify-login must update last_login_at. It is a genuine login event."
    )
    login_ts = user.last_login_at
    if login_ts.tzinfo is None:
        login_ts = login_ts.replace(tzinfo=UTC)
    assert before <= login_ts <= after, (
        f"last_login_at {login_ts!r} should be between {before!r} and {after!r}"
    )


@pytest.mark.asyncio
async def test_token_refresh_does_not_update_last_login_at(
    client: AsyncClient, db_session: AsyncSession
):
    """Token refresh is not a new login. last_login_at must remain unchanged."""
    uid = uuid.uuid4().hex[:8]
    fixed_login = datetime(2026, 2, 10, 8, 0, 0, tzinfo=UTC)
    user = User(
        email=f"synthetic.refreshtest.{uid}@crbcl.ca",
        email_normalized=f"synthetic.refreshtest.{uid}@crbcl.ca",
        password_hash=hash_password("Secret123!"),
        full_name="Refresh Test User",
        is_active=True,
        is_verified=True,
    )
    user.last_login_at = fixed_login
    db_session.add(user)
    await db_session.commit()

    # Login to get a refresh token
    login_res = await client.post(
        "/api/v1/auth/login",
        json={"email": user.email, "password": "Secret123!"},
    )
    assert login_res.status_code == 200
    refresh_token = login_res.json().get("refresh_token")

    # Record last_login_at after login (it was updated by login)
    await db_session.refresh(user)
    after_login_ts = user.last_login_at

    # Perform token refresh
    if refresh_token:
        await client.post(
            "/api/v1/auth/refresh",
            json={"refresh_token": refresh_token},
        )

    # last_login_at must not change during refresh
    await db_session.refresh(user)
    after_refresh_ts = user.last_login_at
    assert after_refresh_ts == after_login_ts, (
        "Token refresh must NOT update last_login_at. "
        "Only successful logins should change this value."
    )
